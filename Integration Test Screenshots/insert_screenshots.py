"""Put iPhone screenshots into the "Actual Results" column of the integration test document.

Folder layout (next to this script):
    pet-owner/      screenshots for the Pet Owner table
    veterinarian/   screenshots for the Veterinarian table

Naming — pick ONE style per folder:
    * Numbered:   01.png, 02.png ... (number = row number in CHECKLIST.txt). Missing numbers stay empty.
    * Unnumbered: leave the iPhone names (IMG_0412.PNG ...). They are used in filename order,
                  so take the screenshots in the same order as the checklist.

Run:   python insert_screenshots.py
Options: --height 3.0   screenshot height in inches (default 3.0)
         --doc  PATH    source document (default: the "- cleaned" copy)
         --out  PATH    output document
"""
import argparse
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
HERE = Path(__file__).resolve().parent
HEADER_ROWS = 6          # title block + column headings at the top of each table
MAX_WIDTH_IN = 2.15      # Actual Results column is ~2.3in wide
IMAGE_EXTS = {'.png', '.jpg', '.jpeg'}
TABLES = [('pet-owner', 0, 'Pet Owner'), ('veterinarian', 1, 'Veterinarian')]


def collect(folder, row_count):
    files = sorted(p for p in folder.iterdir() if p.suffix.lower() in IMAGE_EXTS) if folder.exists() else []
    for p in folder.glob('*') if folder.exists() else []:
        if p.suffix.lower() == '.heic':
            print(f'  ! {p.name}: HEIC is not supported — export it as PNG/JPG first')
    numbered = {int(m.group(1)): p for p in files if (m := re.match(r'(\d+)', p.stem))}
    if numbered:
        extra = [p.name for p in files if not re.match(r'\d+', p.stem)]
        if extra:
            print(f'  ! ignoring files without a row number: {", ".join(extra)}')
        return numbered
    if files and len(files) != row_count:
        print(f'  ! {len(files)} screenshots for {row_count} rows — check the order in the summary below')
    return {i: p for i, p in enumerate(files, 1)}


def insert(tc, table, image, height_in):
    cell = docx.table._Cell(tc, table)
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    para = cell.paragraphs[0]
    for r in list(para._p.iter(W + 'r')):
        r.getparent().remove(r)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(image) as im:
        w, h = im.size
    height = min(height_in, MAX_WIDTH_IN * h / w)   # keep wide images inside the column
    para.add_run().add_picture(str(image), height=Inches(height))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--doc', default=str(HERE.parent / 'Integration Test Document (mobile IOS) - cleaned.docx'))
    ap.add_argument('--out', default=str(HERE.parent / 'Integration Test Document (mobile IOS) - with screenshots.docx'))
    ap.add_argument('--height', type=float, default=3.0)
    args = ap.parse_args()

    d = docx.Document(args.doc)
    total = 0
    for folder, ti, label in TABLES:
        table = d.tables[ti]
        rows = table.rows[HEADER_ROWS:]
        print(f'{label}: {folder}/')
        shots = collect(HERE / folder, len(rows))
        for n, image in sorted(shots.items()):
            if not 1 <= n <= len(rows):
                print(f'  ! {image.name}: row {n} does not exist (table has {len(rows)} rows)')
                continue
            tr = rows[n - 1]._tr
            action = ''.join(t.text or '' for t in tr.findall(W + 'tc')[1].iter(W + 't')).strip()
            insert(tr.findall(W + 'tc')[3], table, image, args.height)
            print(f'  {n:>2}  {image.name:<22} {action[:60]}')
            total += 1
        missing = [n for n in range(1, len(rows) + 1) if n not in shots]
        if missing:
            print(f'  still missing: {", ".join(map(str, missing))}')

    if not total:
        sys.exit('No screenshots found. Put them in the pet-owner/ and veterinarian/ folders first.')
    d.save(args.out)
    print(f'\nInserted {total} screenshots -> {args.out}')


if __name__ == '__main__':
    main()
